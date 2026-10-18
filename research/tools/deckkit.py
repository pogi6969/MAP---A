"""Helpers for building defense-deck slides in the style of the PR2 template decks:
copy a content-slide skeleton (background, crimson header bar, subtitle, kicker) and add
cards, tables, images and equations."""

import copy
import io
import os
import sys

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from insert_video_slides import copy_slide, move_slide  # noqa: E402

CRIMSON = RGBColor(0x8B, 0x07, 0x07)
INK = RGBColor(0x26, 0x26, 0x26)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xC9, 0x8A, 0x10)
CREAM = RGBColor(0xFF, 0xFB, 0xF5)
FONT = "ITC Avant Garde Gothic"
SW, SH = 18288000, 10287000
X0, Y0, CW = 548640, 2697480, 17190720      # content left, top, width
YMAX = 9900000


def runs(shape):
    return [r for p in shape.text_frame.paragraphs for r in p.runs]


def set_text(shape, text, size=None):
    rs = runs(shape)
    if not rs:
        shape.text_frame.text = text
        return
    rs[0].text = text
    for r in rs[1:]:
        r.text = ""
    # drop empty extra paragraphs
    paras = shape.text_frame.paragraphs
    for p in paras[1:]:
        if not "".join(r.text for r in p.runs).strip():
            p._p.getparent().remove(p._p)
    if size:
        rs[0].font.size = Pt(size)


def skeleton(prs, src, title, subtitle, kicker=None):
    """New content slide: copy src (a content slide), keep background, header group,
    title, subtitle and kicker, remove everything else."""
    s = copy_slide(src, prs)
    keep = []
    for sh in list(s.shapes):
        nm = sh.name
        if sh.shape_type in (5, 6) and sh.top <= 0:      # background freeform, header group
            keep.append(sh)
        elif nm in ("Subtitle", "Kicker") or (sh.has_text_frame and sh.top < 100000 and sh.left < 1000000):
            keep.append(sh)
        else:
            sh._element.getparent().remove(sh._element)
    for sh in s.shapes:
        if sh.name == "Subtitle":
            set_text(sh, subtitle)
        elif sh.name == "Kicker":
            if kicker:
                set_text(sh, kicker)
            else:
                sh._element.getparent().remove(sh._element)
        elif sh.has_text_frame and sh.top < 100000:
            n = len(title)
            size = 84 if n <= 13 else 64 if n <= 20 else 50 if n <= 27 else 40
            set_text(sh, title, size)
    return s


def textbox(slide, x, y, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            font=FONT, italic=False):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(60000)
    tf.margin_top = tf.margin_bottom = Emu(30000)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        segs = ln if isinstance(ln, list) else [(ln, bold, color)]
        for seg in segs:
            t, b, c = (seg + (color,))[:3] if len(seg) == 2 else seg
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = b
            r.font.italic = italic
            r.font.name = font
            r.font.color.rgb = c
        p.space_after = Pt(size * 0.35)
    return tb


def card(slide, x, y, w, h, fill=CREAM, line=None, radius=0.06):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(2)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def title_card(slide, x, y, w, h, head, body, head_size=20, body_size=16, fill=CREAM, head_color=CRIMSON, bar=True):
    card(slide, x, y, w, h, fill)
    if bar:
        b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y + 120000), Emu(70000), Emu(h - 240000))
        b.fill.solid()
        b.fill.fore_color.rgb = head_color
        b.line.fill.background()
    textbox(slide, x + 150000, y + 90000, w - 260000, 520000, head, head_size, True, head_color)
    textbox(slide, x + 150000, y + 90000 + int(head_size * 12700 * 1.9), w - 260000,
            h - 200000 - int(head_size * 12700 * 1.9), body, body_size, False, INK)


def table(slide, x, y, w, rows, col_w, size=13, head_size=14, row_h=420000, head_fill=CRIMSON):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Emu(x), Emu(y), Emu(w), Emu(row_h * nr))
    tb = gt.table
    tot = sum(col_w)
    for j, cw in enumerate(col_w):
        tb.columns[j].width = Emu(int(w * cw / tot))
    for i, row in enumerate(rows):
        tb.rows[i].height = Emu(row_h)
        for j, val in enumerate(row):
            cell = tb.cell(i, j)
            cell.margin_left = cell.margin_right = Emu(70000)
            cell.margin_top = cell.margin_bottom = Emu(30000)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.text = ""
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.name = FONT
            if i == 0:
                r.font.size = Pt(head_size)
                r.font.bold = True
                r.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = head_fill
            else:
                r.font.size = Pt(size)
                r.font.color.rgb = INK
                cell.fill.solid()
                cell.fill.fore_color.rgb = CREAM if i % 2 else RGBColor(0xF6, 0xEC, 0xE4)
    return gt


def picture(slide, path, x, y, w=None, h=None):
    from PIL import Image
    im = Image.open(path)
    iw, ih = im.size
    if w and h:
        s = min(w / iw, h / ih)
        pw, ph = int(iw * s), int(ih * s)
        x += (w - pw) // 2
        y += (h - ph) // 2
        return slide.shapes.add_picture(path, Emu(x), Emu(y), Emu(pw), Emu(ph))
    if w:
        return slide.shapes.add_picture(path, Emu(x), Emu(y), width=Emu(w))
    return slide.shapes.add_picture(path, Emu(x), Emu(y), height=Emu(h))


def chain(slide, labels, y, lit_color=GOLD):
    n = len(labels)
    gap = 300000
    w = (CW - gap * (n - 1)) // n
    for i, lab in enumerate(labels):
        x = X0 + i * (w + gap)
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(560000))
        shp.adjustments[0] = 0.5
        shp.fill.solid()
        shp.fill.fore_color.rgb = CRIMSON if i % 2 == 0 else lit_color
        shp.line.fill.background()
        tf = shp.text_frame
        tf.text = ""
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = lab
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.name = FONT
        r.font.color.rgb = WHITE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        if i < n - 1:
            a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(x + w + 40000), Emu(y + 190000), Emu(gap - 80000), Emu(180000))
            a.fill.solid()
            a.fill.fore_color.rgb = INK
            a.line.fill.background()


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def delete_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        if prs.part.related_part(sid.rId) is slide.part:
            prs.part.drop_rel(sid.rId)
            lst.remove(sid)
            return


def index_of(prs, slide):
    for i, s in enumerate(prs.slides):
        if s.part is slide.part:
            return i
    return -1


def latex_dark(tex, path):
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import rationale_video as V
    V.latex_png(tex, path, dpi=300, color="0.10 0.08 0.09")
    return path


def mcs_slide(prs, src, kicker, cards, banner, eqdir, chain_labels=None):
    """The essence of MCS: a chain from open data to proof, then three cards."""
    s = skeleton(prs, src, "HOW MCS COMES INTO PLAY", "The Essence of MCS", kicker)
    chain(slide=s, labels=chain_labels or ["Open data", "Graph", "Algorithm", "Numbers", "Optimisation", "Proof"], y=2620000)
    os.makedirs(eqdir, exist_ok=True)
    gap = 300000
    w = (CW - 2 * gap) // 3
    y = 3400000
    h = 5350000
    for i, (letter, word, body, tex) in enumerate(cards):
        x = X0 + i * (w + gap)
        card(s, x, y, w, h, CREAM)
        textbox(s, x + 180000, y + 60000, 900000, 1150000, letter, 72, True, CRIMSON, font="Times New Roman")
        textbox(s, x + 1000000, y + 330000, w - 1100000, 700000, word, 26, True, CRIMSON)
        textbox(s, x + 180000, y + 1250000, w - 360000, 2500000, body, 19, False, INK)
        if tex:
            png = latex_dark(tex, os.path.join(eqdir, "mcs_%d.png" % i))
            picture(s, png, x + 200000, y + 3850000, w - 400000, 1300000)
    b = card(s, X0, 8950000, CW, 760000, CRIMSON)
    textbox(s, X0, 8950000, CW, 760000, banner, 22, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    return s


def sdg_slide(prs, src, kicker, goals, icon_dir=os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "sdg"), note=None):
    """Real SDG goal icons with the target each study serves."""
    s = skeleton(prs, src, "SUSTAINABLE DEVELOPMENT GOALS", "Goals Served", kicker)
    n = len(goals)
    gap = 300000
    w = (CW - gap * (n - 1)) // n
    icon = min(w - 300000, 3300000)
    for i, (g, head, body) in enumerate(goals):
        x = X0 + i * (w + gap)
        card(s, x, Y0, w, 6900000, CREAM)
        picture(s, os.path.join(icon_dir, "sdg-%02d.png" % g), x + (w - icon) // 2, Y0 + 200000, icon, icon)
        textbox(s, x + 150000, Y0 + 300000 + icon, w - 300000, 600000, head, 20, True, CRIMSON, PP_ALIGN.CENTER)
        textbox(s, x + 150000, Y0 + 950000 + icon, w - 300000, 6900000 - icon - 1100000, body, 15, False, INK, PP_ALIGN.CENTER)
    textbox(s, X0, 9700000, CW, 330000, note or "SDG icons: United Nations. These are intended uses of the model; none has been measured in practice yet.",
            12, False, RGBColor(0x55, 0x4A, 0x4A), PP_ALIGN.CENTER, italic=True)
    return s
