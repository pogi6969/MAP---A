"""Insert the rationale-video opening into a defense deck, the way the team's MAP-A defense
deck does it: after the title slide come a blank background slide, the full-screen
rationale video (plays on click), a tagline slide, and the title slide again. A copy of
the title slide is also placed before QUESTIONS & ANSWER.

Usage:
  python3 insert_video_slides.py DECK.pptx REFERENCE.pptx VIDEO.mp4 POSTER.png OUT.pptx \
      "LINE 1" "LINE 2" "GOLD LINE"
"""

import copy
import io
import subprocess
import sys

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def blank_layout(prs):
    for lay in prs.slide_layouts:
        if lay.name.lower() == "blank":
            return lay
    return prs.slide_layouts[len(prs.slide_layouts) - 1]


def copy_slide(src, dst_prs, skip_media=False):
    """Copy a slide (shapes, background, images) into dst_prs; returns the new slide."""
    new = dst_prs.slides.add_slide(blank_layout(dst_prs))
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rid_map[rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif src.part.package is new.part.package:
            if skip_media and rel.reltype not in (RT.IMAGE,) and "media" in rel.reltype.lower():
                continue
            rid_map[rId] = new.part.relate_to(rel.target_part, rel.reltype)
        elif rel.reltype == RT.IMAGE:
            _, new_rId = new.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
            rid_map[rId] = new_rId
        elif skip_media:
            continue
        else:
            rid_map[rId] = new.part.relate_to(rel.target_part, rel.reltype)
    src_csld = src._element.cSld
    dst_csld = new._element.cSld
    if src_csld.bg is not None:
        if dst_csld.bg is not None:
            dst_csld.remove(dst_csld.bg)
        dst_csld.insert(0, copy.deepcopy(src_csld.bg))
    tree = new.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag in ("nvGrpSpPr", "grpSpPr"):
            continue
        if skip_media and tag == "pic" and el.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}videoFile") is not None:
            continue
        tree.append(copy.deepcopy(el))
    # remap relationship ids in the copied XML (background and shapes)
    for node in new._element.iter():
        for attr in list(node.attrib):
            if attr.startswith("{%s}" % R_NS):
                v = node.attrib[attr]
                if v in rid_map:
                    node.attrib[attr] = rid_map[v]
    return new


def move_slide(prs, slide, index):
    lst = prs.slides._sldIdLst
    target = None
    for sid in list(lst):
        if prs.part.related_part(sid.rId) is slide.part:
            target = sid
            break
    lst.remove(target)
    lst.insert(index, target)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def duration_ms(path):
    out = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", path],
                         capture_output=True, text=True, check=True).stdout.strip()
    return int(round(float(out) * 1000))


def replace_runs(shape, texts):
    runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
    for r, t in zip(runs, texts):
        r.text = t
    for r in runs[len(texts):]:
        r.text = ""


def main():
    deck, ref, video, poster, out, line1, line2, gold = sys.argv[1:9]
    prs = Presentation(deck)
    ex = Presentation(ref)
    ex_blank, ex_video, ex_tag = ex.slides[1], ex.slides[2], ex.slides[3]
    title = prs.slides[0]
    secs = duration_ms(video) / 1000.0

    s_blank = copy_slide(ex_blank, prs)
    set_notes(s_blank, "Blank slide. Keep it up until it is your turn, then click to the rationale video.")

    s_video = copy_slide(ex_video, prs, skip_media=True)
    mv = s_video.shapes.add_movie(video, Emu(0), Emu(0), prs.slide_width, prs.slide_height,
                                  poster_frame_image=poster, mime_type="video/mp4")
    mv.name = "Rationale video"
    # play on the next click, exactly as in the reference deck
    ex_timing = ex_video._element.find(qn("p:timing"))
    if ex_timing is not None:
        tm = copy.deepcopy(ex_timing)
        for el in tm.iter():
            if el.tag == qn("p:spTgt"):
                el.set("spid", str(mv.shape_id))
            if el.tag == qn("p:cTn") and el.get("dur") not in (None, "indefinite", "1"):
                el.set("dur", str(int(round(secs * 1000))))
        old = s_video._element.find(qn("p:timing"))
        if old is not None:
            old.addprevious(tm)
            old.getparent().remove(old)
        else:
            s_video._element.append(tm)
    set_notes(s_video, "Click once to play the rationale video. It runs %d min %02d s. Stay silent and let it play to the title card."
              % (int(secs // 60), int(round(secs % 60))))

    s_tag = copy_slide(ex_tag, prs)
    boxes = [sh for sh in s_tag.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    for sh in boxes:
        txt = sh.text_frame.text
        if "MAP CAN SHOW" in txt or "OUR MODEL" in txt:
            replace_runs(sh, [line1, line2])
            for r in [r for p in sh.text_frame.paragraphs for r in p.runs]:
                r.font.size = Pt(50)
            sh.top = sh.top - Emu(250000)
        elif "COVERING" in txt:
            replace_runs(sh, [gold])
            for r in [r for p in sh.text_frame.paragraphs for r in p.runs]:
                r.font.size = Pt(52)
            sh.left = Emu(0)
            sh.width = prs.slide_width
            sh.top = sh.top + Emu(200000)
    set_notes(s_tag, "Read the two lines on the screen slowly, then introduce the study.")

    s_title2 = copy_slide(title, prs)
    set_notes(s_title2, "Introduce the study title and the researchers, then go to the Introduction.")
    s_title3 = copy_slide(title, prs)
    set_notes(s_title3, "Return to the title before questions.")

    move_slide(prs, s_blank, 1)
    move_slide(prs, s_video, 2)
    move_slide(prs, s_tag, 3)
    move_slide(prs, s_title2, 4)
    # title again before QUESTIONS & ANSWER (or before the last two slides)
    qa = None
    for i, s in enumerate(prs.slides):
        t = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame).upper()
        if "QUESTION" in t and "ANSWER" in t and i > 5:
            qa = i
    if qa is None:
        qa = len(prs.slides) - 3
    move_slide(prs, s_title3, qa)

    cp = prs.core_properties
    cp.author = "De Asis, Dave Justin Ll."
    cp.last_modified_by = "De Asis, Dave Justin Ll."
    cp.comments = ""
    prs.save(out)
    print(out, len(prs.slides), "slides")


if __name__ == "__main__":
    main()
